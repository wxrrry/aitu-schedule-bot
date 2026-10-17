import os
import re
import asyncio
import logging
from typing import Any
import openpyxl

from db.crud import add_group, add_or_update_lesson, normalize_group_name
from schedule_data import get_course_from_group_name

logger = logging.getLogger(__name__)

WEEKDAYS_MAP = {
    "понедельник": 0, "пн": 0,
    "вторник": 1, "вт": 1,
    "среда": 2, "ср": 2,
    "четверг": 3, "чт": 3,
    "пятница": 4, "пт": 4,
    "суббота": 5, "сб": 5
}


def _clean_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _get_sheets_matrix(file_path: str) -> list[tuple[str, list[list[str]]]]:
    sheets_matrix = []
    if file_path.lower().endswith(".xls"):
        try:
            import xlrd
            book = xlrd.open_workbook(file_path)
            for sheet in book.sheets():
                matrix = []
                for r in range(min(sheet.nrows, 350)):
                    row = [_clean_str(sheet.cell_value(r, c)) for c in range(min(sheet.ncols, 60))]
                    matrix.append(row)
                sheets_matrix.append((sheet.name, matrix))
            return sheets_matrix
        except Exception as e:
            logger.warning(f"Ошибка чтения через xlrd: {e}")

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            matrix = []
            max_row = min(sheet.max_row, 350)
            max_col = min(sheet.max_column, 60)
            for r in range(1, max_row + 1):
                row = [_clean_str(sheet.cell(row=r, column=c).value) for c in range(1, max_col + 1)]
                matrix.append(row)
            sheets_matrix.append((sheet_name, matrix))
    except Exception as e:
        logger.error(f"Ошибка чтения через openpyxl: {e}")

    return sheets_matrix


def parse_excel_sync(file_path: str) -> dict[str, Any]:
    """
    Синхронный парсинг Excel-файла семестрового расписания (.xlsx / .xls).
    Поддерживает стандартные структуры расписания СПО/колледжа:
    - По колонкам групп
    - Строки дней недели, номеров пар (1..6), числителя/знаменателя
    """
    results = {
        "groups_found": set(),
        "lessons": [],
        "lessons_parsed": 0,
        "errors": []
    }

    group_regex = re.compile(r'\b(\d{1,2}[А-Яа-яA-Za-z]{1,4}-\d{1,3}[а-яА-Яa-zA-Z]?)\b')
    room_regex = re.compile(r'(?:каб\.?|ауд\.?|№)?\s*(\b\d{2,4}[а-яА-Я]?\b|спортзал|актовый\s*зал|стадион|тир|дистант)', re.IGNORECASE)
    teacher_regex = re.compile(r'\b([А-ЯЁ][а-яё]+)\s+([А-ЯЁ]\.\s*[А-ЯЁ]\.?)\b')

    sheets = _get_sheets_matrix(file_path)
    for sheet_name, matrix in sheets:
        if not matrix:
            continue
        max_row = len(matrix)

        # 1. Поиск колонок или ячеек с названиями групп
        group_cols: dict[int, str] = {}
        for r in range(min(max_row, 10)):
            row = matrix[r]
            for c in range(len(row)):
                val = row[c]
                matches = group_regex.findall(val)
                for m in matches:
                    grp_name = normalize_group_name(m)
                    group_cols[c] = grp_name
                    results["groups_found"].add(grp_name)

        if not group_cols:
            continue

        current_weekday = None
        current_pair_num = None

        for r in range(max_row):
            row = matrix[r]
            if not row:
                continue
            row_text = " ".join(row[:min(len(row), 6)]).lower()

            # Определение дня недели
            for day_name, day_idx in WEEKDAYS_MAP.items():
                if day_name in row_text:
                    current_weekday = day_idx
                    break

            # Определение номера пары (1..6)
            for c in range(min(len(row), 5)):
                cell_val = row[c]
                if cell_val.isdigit() and 1 <= int(cell_val) <= 6:
                    current_pair_num = int(cell_val)
                    break

            if current_weekday is None or current_pair_num is None:
                continue

            # Определение четности недели на уровне строки
            row_parity = "both"
            for c in range(min(len(row), 5)):
                val_lower = row[c].lower().strip()
                if any(k in val_lower for k in ["числ", "1 нед", "i нед"]) or val_lower == "ч":
                    row_parity = "chislitel"
                    break
                elif any(k in val_lower for k in ["знам", "2 нед", "ii нед"]) or val_lower == "з":
                    row_parity = "znamenatel"
                    break

            # Считывание предметов по найденным колонкам групп
            for col_idx, grp_name in group_cols.items():
                if col_idx >= len(row):
                    continue
                cell_content = row[col_idx]
                if not cell_content or len(cell_content) < 3:
                    continue

                # Извлечение кабинета
                room = "—"
                room_m = room_regex.search(cell_content)
                if room_m:
                    room = room_m.group(1).strip()

                # Извлечение преподавателя
                teacher = "—"
                t_m = teacher_regex.search(cell_content)
                if t_m:
                    teacher = f"{t_m.group(1)} {t_m.group(2)}".strip()

                # Очистка названия предмета
                subj_clean = cell_content
                if t_m:
                    subj_clean = subj_clean.replace(t_m.group(0), "")
                if room_m and room_m.group(0).lower() != "дистант":
                    subj_clean = subj_clean.replace(room_m.group(0), "")
                subj_clean = re.sub(r'[\(\)\[\]]', ' ', subj_clean)
                subj_clean = re.sub(r'\s+', ' ', subj_clean).strip()
                if not subj_clean or len(subj_clean) < 2:
                    subj_clean = cell_content

                if row_parity != "both":
                    results["lessons"].append({
                        "group_name": grp_name,
                        "weekday": current_weekday,
                        "pair_num": current_pair_num,
                        "parity": row_parity,
                        "subject": subj_clean,
                        "room": room,
                        "teacher": teacher
                    })
                    results["lessons_parsed"] += 1
                elif "/" in cell_content and len(cell_content.split("/")) == 2:
                    parts = cell_content.split("/")
                    for parity, subj_part in [("chislitel", parts[0].strip()), ("znamenatel", parts[1].strip())]:
                        if subj_part and len(subj_part) > 1 and subj_part not in ("-", "—"):
                            results["lessons"].append({
                                "group_name": grp_name,
                                "weekday": current_weekday,
                                "pair_num": current_pair_num,
                                "parity": parity,
                                "subject": subj_part,
                                "room": room,
                                "teacher": teacher
                            })
                            results["lessons_parsed"] += 1
                else:
                    results["lessons"].append({
                        "group_name": grp_name,
                        "weekday": current_weekday,
                        "pair_num": current_pair_num,
                        "parity": "both",
                        "subject": subj_clean,
                        "room": room,
                        "teacher": teacher
                    })
                    results["lessons_parsed"] += 1

    return results


async def parse_and_apply_excel_schedule(file_path: str) -> dict[str, Any]:
    """
    Асинхронная обёртка для парсинга и сохранения расписания в SQLite.
    Выполняется в отдельном потоке (asyncio.to_thread), чтобы не блокировать бота.
    """
    results = await asyncio.to_thread(parse_excel_sync, file_path)
    
    # Сохраняем найденные группы в БД
    for grp in results.get("groups_found", []):
        course = get_course_from_group_name(grp)
        await add_group(grp, course=course)

    # Сохраняем все уроки в БД
    applied_count = 0
    for lesson in results.get("lessons", []):
        await add_or_update_lesson(
            group_name=lesson["group_name"],
            weekday=lesson["weekday"],
            pair_num=lesson["pair_num"],
            parity=lesson["parity"],
            subject=lesson["subject"],
            room=lesson.get("room", "—"),
            teacher=lesson.get("teacher", "—")
        )
        applied_count += 1

    results["applied_count"] = applied_count
    return results
